"""Bobur's daily «Аудиозаписи (<Месяц>-26)» file, filled automatically.
Every day: 6 ТП from different regions (not checked yet this month), 2 pharmacy visits each;
Аптека, Длительность записи, Комментария (AI listened, in Bobur's wording) and Ссылка are filled.
Newest day on top, same layout and styles as his file. The file lives encrypted in the store (az_file)."""
import copy, datetime as dt, io, random, re, tempfile
from zoneinfo import ZoneInfo
import pandas as pd
from openpyxl import load_workbook

import store
import audio
from sales import BASE

TZ = ZoneInfo("Asia/Tashkent")
PER_DAY, PER_TP = 6, 2
SHORT = {'бух': 'Бухара', 'джиз': 'Джизак', 'дж': 'Джизак', 'нав': 'Наваи', 'нук': 'Нукус', 'сам': 'Самарканд',
         'тер.в': 'Термез', 'тер.з': 'Термез', 'тер': 'Термез', 'нам': 'Наманган', 'ург': 'Ургенч', 'фер': 'Фергана',
         'кок': 'Коканд', 'кар': 'Карши'}            # Bobur's file covers these regions only (no Ташкент)
MON = {1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель', 5: 'Май', 6: 'Июнь', 7: 'Июль', 8: 'Август',
       9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь'}

PROMPT = """This is an audio recording made by a pharmaceutical sales representative (ТП) during a visit to a pharmacy in
Uzbekistan (Uzbek or Russian speech). Write the manager's check comment in Uzbek Cyrillic, in exactly this style:
- "Маълумот йўқ" - nothing useful: silence, noise, no work conversation, or the rep does not talk with the pharmacist about work
- "Остатка олинди" - the rep asked / counted the remaining stock (остатка, қолдиқ, нечта қолди) of our products
- "Анализ қилинди товарлар сотилиши бўйича. Остатка олинди" - they discussed how our products sell AND stock was taken
- otherwise a very short phrase of what was done, e.g. "Презентация қилинди (Риномакс)", "Мерч ҳақида гапирилди", "Заявка олинди"
  (several may be joined with ". ")
Our products: Риномакс, Форсил, Риноксил, Мирамистин, Эрегра, Тивомакс, Триомакс, Альцетро, Авирол.
Return ONLY JSON {"comment": "..."}"""


def short_region(name: str):
    m = re.match(r'^\(([^)]*)\)', str(name).strip())
    return SHORT.get(m.group(1).strip().lower()) if m else None


def load_file():
    data = store.load_bytes('az_file')
    if not data:                             # first run: Bobur's file, encrypted with the roster key
        import os, roster as R
        p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'az_base.enc')
        data = R._fernet().decrypt(open(p, 'rb').read()) if os.path.exists(p) else None
    return load_workbook(io.BytesIO(data)) if data else None


def save_file(wb):
    b = io.BytesIO(); wb.save(b)
    store.save_bytes('az_file', b.getvalue())


def blocks(ws):
    """[(day 'dd.mm', [manager names])] from the file, newest first."""
    out, cur = [], None
    for r in ws.iter_rows(min_row=1, max_col=2):
        if r[1].value == 'Менеджер':
            cur = (str(r[0].value).strip(), []); out.append(cur)
        elif cur and r[1].value:
            cur[1].append(str(r[1].value).strip())
    return out


def key(name):
    import roster as R
    return R.key(name)


def pick(day: dt.date, visits: pd.DataFrame, checked: set, last_seen: dict, seed: int):
    """6 ТП from different regions who visited pharmacies that day; first those not checked this month."""
    rnd = random.Random(seed)
    v = visits[visits['day'].dt.date == day].copy()
    v = v[v['Фарм. пред'].str.contains('Торг', na=False)]
    v['who'] = v['Фарм. пред'].str.replace(r'\s*\([^()]*\)\s*$', '', regex=True).str.strip()
    v['reg'] = v['who'].map(short_region)
    v = v[v['reg'].notna()]
    cnt = v.groupby('who').size()
    people = [w for w in cnt.index if cnt[w] >= PER_TP]
    fresh = [w for w in people if key(w) not in checked]
    old = sorted([w for w in people if key(w) in checked], key=lambda w: last_seen.get(key(w), ''))
    rnd.shuffle(fresh)
    chosen, regs = [], []
    for pool in (fresh, old):                 # different regions first, then fill up
        for w in pool:
            r = v[v.who == w]['reg'].iloc[0]
            if len(chosen) < PER_DAY and w not in chosen and r not in regs:
                chosen.append(w); regs.append(r)
    for pool in (fresh, old):
        for w in pool:
            if len(chosen) < PER_DAY and w not in chosen:
                chosen.append(w)
    rows = []
    for w in sorted(chosen, key=lambda x: (v[v.who == x]['reg'].iloc[0], x)):
        vv = v[v.who == w]
        ids = rnd.sample(list(vv['#'].astype(int)), PER_TP)
        for i in sorted(ids):
            r = vv[vv['#'].astype(int) == i].iloc[0]
            rows.append({'who': w, 'reg': r['reg'], 'id': i,
                         'obj': re.sub(r'\s*\([^()]*\)\s*$', '', str(r['Контр-агент/Врач'])).strip()})
    return rows


def listen(s, rows):
    for r in rows:
        r['dur'], r['comment'] = None, ''
        try:
            url = audio.audio_url(s, r['id'])
            if not url:
                r['comment'] = 'Запись йўқ'; continue
            ext = url.split('?')[0].rsplit('.', 1)[-1].lower()
            data = s.get(url, timeout=120).content
            with tempfile.NamedTemporaryFile(suffix='.' + ext) as f:
                f.write(data); f.flush()
                m = audio.measure(f.name)
            r['dur'] = m['dur']
            if m['dur'] < 5 or m['silence_pct'] >= 97:
                r['comment'] = 'Маълумот йўқ (ёзувда овоз йўқ)'; continue
            a = audio.ai(data, audio.MIME.get(ext, 'audio/aac'), PROMPT)
            r['comment'] = str(a.get('comment') or 'Маълумот йўқ').strip()
        except Exception as e:
            print("::warning::az listen " + str(e)[:120])
            r['comment'] = r['comment'] or 'Текширилмади (ИИ лимити)'
    return rows


def refill(ws) -> int:
    """Rows left as 'Текширилмади (ИИ лимити)' are listened to again."""
    from sales import login, switch_project
    todo = []
    for row in ws.iter_rows(min_row=1):
        g, h = row[6].value, row[7].value
        if isinstance(g, str) and g.startswith('Текширилмади') and isinstance(h, str):
            m = re.search(r'/visits/(\d+)/', h)
            if m:
                todo.append((row, {'id': int(m.group(1))}))
    if not todo:
        return 0
    s = login(); switch_project(s, 1)
    listen(s, [r for _, r in todo])
    n = 0
    for row, r in todo:
        if r['comment'] and not r['comment'].startswith('Текширилмади'):
            row[6].value = r['comment']; n += 1
            if r.get('dur'):
                sec = int(round(r['dur'])); row[5].value = dt.time(min(sec // 60, 23), sec % 60)
    print(f"::notice::az refill {n}/{len(todo)}")
    return n


def write_block(ws, day: dt.date, rows):
    """Insert a day block at the top, copying the styles of the existing header/data rows."""
    hdr = [copy.copy(ws.cell(1, j)._style) for j in range(1, 9)]
    dat = [copy.copy(ws.cell(2, j)._style) for j in range(1, 9)]
    hh, dh = ws.row_dimensions[1].height, ws.row_dimensions[2].height
    n = len(rows) + 1
    ws.insert_rows(1, n)
    for j, v in enumerate([f"{day:%d.%m}", 'Менеджер', 'Регион', 'Дата', 'Аптека', 'Длительность записи', 'Комментария', 'Ссылка'], 1):
        c = ws.cell(1, j, v); c._style = copy.copy(hdr[j - 1])
    ws.row_dimensions[1].height = hh
    for i, r in enumerate(rows, 2):
        dur = None
        if r['dur']:
            sec = int(round(r['dur']))
            dur = dt.time(min(sec // 60, 23), sec % 60)        # his file keeps m:ss in an h:mm cell
        link = f"{BASE}/lists/visits/{r['id']}/result"
        vals = [None, r['who'], r['reg'], f"{day:%d.%m}", f'"{r["obj"]}"', dur, r['comment'], link]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c._style = copy.copy(dat[j - 1])
        ws.row_dimensions[i].height = dh
    # openpyxl does not move hyperlinks on insert_rows: rebuild them from the link texts
    ws._hyperlinks = []
    for row in ws.iter_rows(min_col=8, max_col=8):
        c = row[0]
        if isinstance(c.value, str) and c.value.startswith('http'):
            c.hyperlink = c.value


def run(days, send=True):
    """Add blocks for the given days (oldest first, so the newest ends on top) and send the file."""
    import vcache
    from sales import login, switch_project
    wb = load_file()
    if wb is None:
        print("::warning::az: base file not uploaded yet"); return
    ws = wb.active
    refilled = refill(ws)
    have = {b[0] for b in blocks(ws)}
    days = [d for d in sorted(days) if f"{d:%d.%m}" not in have]
    if not days and not refilled:
        print("::notice::az: nothing new"); return
    v = vcache.get('tp', min(days).replace(day=1), max(days))
    v = v[v['Тип визита'].fillna('').str.contains('аптек|контрагент', case=False)]
    v = v.assign(day=pd.to_datetime(v['Дата'], format='%d.%m.%Y'))
    s = login(); switch_project(s, 1)
    added = []
    for d in days:
        checked, last_seen = set(), {}
        for b in blocks(ws):
            try:
                bd = dt.date(d.year, int(b[0][3:5]), int(b[0][:2]))
            except Exception:
                continue
            for m in b[1]:
                last_seen[key(m)] = max(last_seen.get(key(m), ''), bd.isoformat())
                if bd.month == d.month and bd.year == d.year:
                    checked.add(key(m))
        rows = pick(d, v, checked, last_seen, seed=d.toordinal())
        if not rows:
            print(f"::notice::az {d}: no ТП visits"); continue
        listen(s, rows)
        write_block(ws, d, rows)
        added.append((d, rows))
        print(f"::notice::az {d}: {len(rows) // PER_TP} ТП, {len(rows)} rows")
    save_file(wb)
    if send and (added or refilled):
        from telegram import send_text, send_file
        last = max([d for d, _ in added] or [dt.datetime.now(TZ).date()])
        path = f"/tmp/Аудиозаписи ({MON[last.month]}-{last:%y}).xlsx"
        wb.save(path)
        lines = [f"🎧 <b>Аудиозаписи</b>" + (" — қўшилди: " + ", ".join(f"{d:%d.%m}" for d, _ in added) if added else "")]
        if refilled:
            lines.append(f"• олдин ИИ лимити туфайли текширилмай қолган {refilled} та ёзув тингланди")
        for d, rows in added:
            bad = sum(1 for r in rows if r['comment'].startswith(('Маълумот йўқ', 'Запись йўқ')))
            lines.append(f"• {d:%d.%m}: {len(rows) // PER_TP} ТП, {len(rows)} ёзув, «Маълумот йўқ / ёзув йўқ»: {bad}")
        send_text("\n".join(lines))
        send_file(path, "Аудиозаписи — янги кунлар тепада")


if __name__ == "__main__":
    import os, sys
    mode = os.environ.get("AZ_MODE", "")
    today = dt.datetime.now(TZ).date()
    if mode == "upload":                     # Bobur's own file -> becomes the base
        src = sys.argv[1]
        store.save_bytes('az_file', open(src, 'rb').read())
        print("::notice::az base saved")
    else:
        d1 = dt.date.fromisoformat(os.environ["AZ_FROM"]) if os.environ.get("AZ_FROM") else today - dt.timedelta(days=1)
        d2 = dt.date.fromisoformat(os.environ["AZ_TO"]) if os.environ.get("AZ_TO") else today - dt.timedelta(days=1)
        run([d1 + dt.timedelta(days=i) for i in range((d2 - d1).days + 1)])
