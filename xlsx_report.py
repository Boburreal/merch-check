"""Build the Excel report (same layout as the yearly file) for any period."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from analysis import FLAGS, GAP_MIN

MON = {1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель', 5: 'Май', 6: 'Июнь', 7: 'Июль', 8: 'Август',
       9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь'}
HDR = PatternFill('solid', fgColor='1F7A3A'); HF = Font(bold=True, color='FFFFFF')
TOT = PatternFill('solid', fgColor='C6EFCE'); RED = PatternFill('solid', fgColor='FFC7CE')
CAT = {'merch': 'Мерч бор', 'none': 'Мерч йўқ', 'suspicious': 'Шубҳали (яқин план)', 'no_photo': 'Фото йўқ',
       'error': 'Хато (қайта текширилади)'}
CATFILL = {'merch': 'C6EFCE', 'none': 'FFC7CE', 'suspicious': 'FFEB9C', 'no_photo': 'D9D9D9', 'error': 'D9D9D9'}


def _sheet(wb, name, df, total=True, pct=(), dec=(), widths=None, links=()):
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, 'item') else v) for v in r])
    n = ws.max_row
    if total and n > 1:
        ws.append(['ЖАМИ'] + [None] * (len(df.columns) - 1))
        for j, c in enumerate(df.columns, 1):
            if j > 1 and pd.api.types.is_numeric_dtype(df[c]) and c not in pct and c not in dec:
                L = get_column_letter(j)
                ws.cell(n + 1, j, f"=SUM({L}2:{L}{n})")
        for c in ws[n + 1]:
            c.font = Font(bold=True); c.fill = TOT
    for c in ws[1]:
        c.font = HF; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 36
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, 12 if j > 1 else 34)
        fmt = '0.0%' if c in pct else ('0.0' if c in dec else None)
        if fmt:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, j).number_format = fmt
        if c in links:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, j); u = cell.value
                if isinstance(u, str) and u.startswith('http'):
                    cell.hyperlink = u; cell.value = 'Очиш'; cell.font = Font(color='0563C1', underline='single')
    ws.freeze_panes = 'B2'
    if not total:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _summary(t, who, merch=None):
    gr = t.groupby(who)
    sv = pd.DataFrame({
        who: gr.size().index,
        'Регион': gr['Регион'].first().values,
        'Лавозим': gr['Лавозим'].agg(lambda s: s.mode().iat[0]).values,
        'Иш кунлари': gr['day'].nunique().values,
        'Факт визит': gr.size().values,
        'Уникальный объект': gr['obj'].nunique().values,
    })
    sv['Кунига ўртача'] = (sv['Факт визит'] / sv['Иш кунлари']).round(1)
    for k, n in FLAGS.items():
        sv[n] = gr['f_' + k].sum().values
    sv['Шубҳали визит (жами)'] = gr['Шубҳали'].sum().values
    sv['Шубҳали %'] = sv['Шубҳали визит (жами)'] / sv['Факт визит']
    if merch is not None:
        m = t[[who, 'id']].merge(merch, on='id', how='left')
        chk = m['category'].isin(['merch', 'none', 'suspicious'])
        mg = m.assign(chk=chk, mb=m['category'].eq('merch'), sus=m['category'].eq('suspicious'),
                      v=pd.to_numeric(m['vitrina'], errors='coerce').fillna(0),
                      k=pd.to_numeric(m['kassa'], errors='coerce').fillna(0)).groupby(who)
        sv['Мерч текширилди'] = mg['chk'].sum().reindex(sv[who]).values
        sv['Мерч бор'] = mg['mb'].sum().reindex(sv[who]).values
        sv['Витрина'] = mg['v'].sum().reindex(sv[who]).values.astype(int)
        sv['Касса'] = mg['k'].sum().reindex(sv[who]).values.astype(int)
        sv['Мерч шубҳали (яқин план)'] = mg['sus'].sum().reindex(sv[who]).values
        sv['Мерч %'] = np.where(sv['Мерч текширилди'] > 0, sv['Мерч бор'] / sv['Мерч текширилди'].replace(0, np.nan), np.nan)
    return sv.sort_values(who)


def _pivot(t, who, by_month):
    if by_month:
        key = t.day.dt.month.map(MON); order = [MON[m] for m in sorted(t.day.dt.month.unique())]
    else:
        key = t.day.dt.strftime('%d.%m'); order = sorted(key.unique(), key=lambda s: (s[3:], s[:2]))
    pv = t.assign(k=key).groupby([who, 'k']).size().unstack('k').reindex(columns=order)
    pv['Жами'] = pv.sum(axis=1)
    return pv.reset_index()


def _list(t, who, cols_extra):
    sl = t[t['Шубҳали']][['#', 'Дата', 'Время завершения визита', who, 'Контр-агент/Врач', 'Адрес', 'Изоҳ', 'Фото 1', 'Фото 2']].copy()
    sl.columns = ['ID', 'Дата', 'Ёпилган вақт', who, 'Объект', 'Адрес', 'Изоҳ (нима учун шубҳали)', 'Фото 1', 'Фото 2']
    return sl


W_LIST = {'ID': 9, 'Дата': 11, 'Ёпилган вақт': 16, 'ТП': 32, 'МП': 32, 'Объект': 40, 'Адрес': 30,
          'Изоҳ (нима учун шубҳали)': 70, 'ИИ изоҳи (ҳар бир фото)': 75, 'Хулоса': 18, 'Фото 1': 8, 'Фото 2': 8}


def build(path, title, tp, mp, merch, accuracy_note, loc_changes=None, loc_static=None):
    wb = Workbook()
    ws = wb.active; ws.title = 'Изоҳ'
    days = (tp.day.max() - tp.day.min()).days + 1 if len(tp) else 1
    by_month = days > 31
    lines = [
        (title, True), ('', False),
        ('Кўрсаткичлар', True),
        ('Факт — ёпилган визитлар сони. Уникальный — шу даврда борилган турли объектлар сони.', False),
        ('Шубҳали — қуйидаги белгилардан камида биттаси бор визит. Бу «сохта» дегани эмас, текширишга арзийди дегани:', False),
        (f'  • Тез ёпилган: шу куни олдинги визит (бошқа объект) дан {GAP_MIN} дақиқа ёки камроқ вақт ўтиб ёпилган.', False),
        ('  • Локацияси ўзгарган аптека: визитдан олдинги 3 кун ичида аптека локацияси ўчирилган/кўчирилган/қайта қўйилган • Фото йўқ • Фото 1 = Фото 2 • Фото бошқа визитда қайта ишлатилган', False),
        ('  • Такрорий визит: шу куни шу объектга иккинчи марта • Иш вақтидан ташқари: 08:00 дан олдин ёки 20:00 дан кейин.', False),
        ('', False),
        ('Мерч (фақат ТП)', True),
        ('Қоида: Риномакс ХОТ ва Форсил иккаласи БИТТА расмда. Витрина — витринада/полкада; Касса — касса олдида.', False),
        ('Яқин масофадан олинган, аптека кўринмайдиган расм — «Шубҳали (яқин план)», мерчга қўшилмайди.', False),
        (accuracy_note, False),
        ('«ТП мерч рўйхат» варағида ҳар бир визит бўйича ИИ ҳар бир фотода нимани кўргани ва фото линклари бор.', False),
        ('', False),
        ('Локация назорати', True),
        ('Sales аптека локацияси тарихини сақламайди. Шунинг учун ҳар соатда аптекалар рўйхатининг «сурати» олинади ва олдингиси билан солиштирилади.', False),
        ('«Локация ўзгаришлари» — қайси аптеканинг локацияси қачон ўчирилгани, қўйилгани ёки неча км га кўчирилгани (вақт — иккита сурат орасидаги оралиқ).', False),
        ('«Локация шубҳали аптекалар» — шу даврда визит бўлган, координатаси шубҳали аптекалар: 3+ аптека бир нуқтада (~10 м), Ўзбекистондан ташқарида, ўз вилоятидан 100+ км узоқда.', False),
    ]
    for txt, b in lines:
        ws.append([txt]); c = ws.cell(ws.max_row, 1)
        c.font = Font(bold=b, size=13 if b else 11); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 140

    if len(tp):
        sv = _summary(tp, 'ТП', merch)
        s = _sheet(wb, 'ТП свод', sv, pct=('Шубҳали %', 'Мерч %'), dec=('Кунига ўртача',))
        col = list(sv.columns).index('Шубҳали %') + 1
        for r in range(2, s.max_row):
            v = s.cell(r, col).value
            if v is not None and v >= 0.4:
                s.cell(r, col).fill = RED
        if days > 1:
            _sheet(wb, 'ТП ойма-ой' if by_month else 'ТП кунлик', _pivot(tp, 'ТП', by_month))
        if merch is not None and len(merch):
            ml = tp[['id', '#', 'Дата', 'Время завершения визита', 'ТП', 'Контр-агент/Врач', 'Фото 1', 'Фото 2']].merge(merch, on='id', how='inner')
            ml = pd.DataFrame({
                'ID': ml['#'], 'Дата': ml['Дата'], 'Ёпилган вақт': ml['Время завершения визита'], 'ТП': ml['ТП'],
                'Объект': ml['Контр-агент/Врач'], 'Хулоса': ml['category'].map(CAT).fillna(ml['category']),
                'Витрина': np.where(ml['vitrina'].astype(str) == '1', 'Да', 'Нет'),
                'Касса': np.where(ml['kassa'].astype(str) == '1', 'Да', 'Нет'),
                'ИИ изоҳи (ҳар бир фото)': ml['comment'].fillna(''), 'Фото 1': ml['Фото 1'], 'Фото 2': ml['Фото 2'],
                '_cat': ml['category']}).sort_values(['ТП', 'Дата', 'Ёпилган вақт'])
            cats = ml.pop('_cat').tolist()
            s = _sheet(wb, 'ТП мерч рўйхат', ml, total=False, widths=W_LIST, links=('Фото 1', 'Фото 2'))
            for r, c in enumerate(cats, 2):
                s.cell(r, 6).fill = PatternFill('solid', fgColor=CATFILL.get(c, 'FFFFFF'))
        _sheet(wb, 'ТП шубҳали рўйхат', _list(tp, 'ТП', ()), total=False, widths=W_LIST, links=('Фото 1', 'Фото 2'))
    if loc_changes is not None and len(loc_changes):
        _sheet(wb, 'Локация ўзгаришлари', loc_changes, total=False,
               widths={'Вақт оралиғи': 34, 'Аниқланди': 16, 'ID': 8, 'Аптека': 40, 'Вилоят': 16, 'Туман': 18,
                       'Бириктирилган ТП': 45, 'Ўзгариш': 20, 'Масофа (км)': 11, 'Эски координата': 22, 'Янги координата': 22})
    if loc_static is not None and len(loc_static):
        _sheet(wb, 'Локация шубҳали аптекалар', loc_static, total=False, links=('Харита',),
               widths={'ID': 8, 'Аптека': 40, 'Вилоят': 16, 'Туман': 18, 'Бириктирилган ТП': 45, 'Координата': 22,
                       'Изоҳ (нима учун шубҳали)': 60, 'Визитлар (шу давр)': 12, 'Харита': 9})
    if mp is not None and len(mp):
        _sheet(wb, 'МП свод', _summary(mp, 'МП'), pct=('Шубҳали %',), dec=('Кунига ўртача',))
        if days > 1:
            _sheet(wb, 'МП ойма-ой' if by_month else 'МП кунлик', _pivot(mp, 'МП', by_month))
        _sheet(wb, 'МП шубҳали рўйхат', _list(mp, 'МП', ()), total=False, widths=W_LIST, links=('Фото 1', 'Фото 2'))
    wb.save(path)
