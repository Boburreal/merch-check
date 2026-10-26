"""Monthly ТП file in Bobur's "Факт визит и мерч" format (values only, no formulas - Protected View safe).
Sheets: Ежедневный (day blocks + ИТОГО), Ежемесячный (per ТП + per region), Визитлар тўлиқ (every visit),
Аптека муаммолари (pharmacy-level issues), Локация ўзгаришлари, Изоҳ (rules)."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import roster as R
from analysis import FLAGS

MON = {1: 'Январь', 2: 'Февраль', 3: 'Март', 4: 'Апрель', 5: 'Май', 6: 'Июнь', 7: 'Июль', 8: 'Август',
       9: 'Сентябрь', 10: 'Октябрь', 11: 'Ноябрь', 12: 'Декабрь'}
HDR = PatternFill('solid', fgColor='1F7A3A'); HF = Font(bold=True, color='FFFFFF')
GRN = PatternFill('solid', fgColor='92D050'); RED = PatternFill('solid', fgColor='FFC7CE'); YEL = PatternFill('solid', fgColor='FFEB9C')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CAT = {'merch': 'Мерч бор', 'none': 'Мерч йўқ', 'suspicious': 'Шубҳали (яқин план)', 'no_photo': 'Фото йўқ',
       'error': 'Текширилмади'}
SHORT = {'gap': 'тез ёпилган', 'loc': 'локацияси ўзгарган аптека', 'photo': 'фото йўқ', 'same': 'фото1=фото2',
         'reuse': 'фото қайта ишлатилган', 'dup': 'такрорий визит', 'hours': 'иш вақтидан ташқари',
         'region_out': 'бошқа ҳудуддаги аптека', 'cluster': 'бир нуқтада 5+ аптека', 'multi': 'аптекага 3+ ТП',
         'over7': 'бир аптекага 7+ визит', 'closeup': 'мерч фото яқин план'}


def _hdr(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = HF; x.fill = HDR; x.border = BOX
        x.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 32


def _pct(a, b):
    return round(a / b, 4) if b else 0


def day_comment(g: pd.DataFrame, unchecked: int) -> str:
    parts = []
    for k, n in SHORT.items():
        c = int(g['f_' + k].sum()) if 'f_' + k in g else 0
        if c:
            parts.append(f"{c} та {n}")
    if unchecked:
        parts.append(f"{unchecked} та визит мерчи ҳали текширилмади")
    new = g['Янги'].any() if 'Янги' in g else False
    if new:
        parts.insert(0, "ЯНГИ ТП")
    return "; ".join(parts)


def rows_per_day(tp: pd.DataFrame, merch: pd.DataFrame, seed: pd.DataFrame, plans: dict):
    """One row per ТП per working day: plan (1 day), fact, vitrina, kassa, comment."""
    m = tp[['id']].merge(merch, on='id', how='left') if len(merch) else tp[['id']].assign(category=None, vitrina=0, kassa=0)
    tp = tp.assign(cat=m['category'].values,
                   v=pd.to_numeric(m.get('vitrina'), errors='coerce').fillna(0).values,
                   k=pd.to_numeric(m.get('kassa'), errors='coerce').fillna(0).values)
    tp['f_closeup'] = tp['cat'].eq('suspicious')
    seed_map = {}
    if seed is not None and len(seed):
        for r in seed.itertuples(index=False):
            seed_map[(r.date, r.key)] = (r.vitrina, r.kassa)
    out = []
    for (day, name), g in tp.groupby(['day', 'ТП'], sort=True):
        k = R.key(name)
        ds = day.strftime('%Y-%m-%d')
        checked = g['cat'].isin(['merch', 'none', 'suspicious']).sum()
        vit, kas, src = int(g.v.sum()), int(g.k.sum()), 'ИИ'
        if checked == 0 and (ds, k) in seed_map:
            vit, kas = [int(x) if pd.notna(x) else 0 for x in seed_map[(ds, k)]]
            src = 'қўлда'
        out.append({'day': day, 'ТП': name, 'Должность': 'ТП', 'Регион': g['Иш ҳудуди'].iloc[0],
                    'План': plans.get(k, R.DEFAULT_PLAN), 'Факт': len(g), 'Витрина': vit, 'Касса': kas,
                    'src': src, 'Комментария': day_comment(g, len(g) - checked if src == 'ИИ' else 0)})
    return pd.DataFrame(out), tp


def build(path, month_title, tp, merch, seed, people, loc_changes):
    plans = dict(zip(people['key'], people['plan']))
    days, tpx = rows_per_day(tp, merch, seed, plans)
    wb = Workbook()

    # ---------- Ежедневный ----------
    ws = wb.active; ws.title = 'Ежедневный'
    cols = ['Дата', '№', 'Менеджер', 'Должность', 'Регион', 'Кол-во дней', 'План Sales', 'Визит факт', '%Выполнения',
            'Витрина Мерч', '%Выполнения', 'Касса Мерч', '%Выполнения', 'Комментария']
    _hdr(ws, 1, cols)
    r = 2; totals = []
    for day, g in days.groupby('day', sort=True):
        g = g.sort_values('ТП', key=lambda s: s.str.lower())
        for n, x in enumerate(g.itertuples(index=False), 1):
            vals = [day.strftime('%d.%m'), n, x.ТП, 'ТП', x.Регион, 1, x.План, x.Факт, _pct(x.Факт, x.План),
                    x.Витрина, _pct(x.Витрина, x.План), x.Касса, _pct(x.Касса, x.План),
                    (x.Комментария + ("; мерч қўлда киритилган" if x.src == 'қўлда' else '')).strip('; ')]
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, v); c.border = BOX
                if j in (9, 11, 13): c.number_format = '0%'
            if x.Комментария:
                ws.cell(r, 14).fill = YEL
            r += 1
        s = g[['План', 'Факт', 'Витрина', 'Касса']].sum()
        vals = [day.strftime('%d.%m'), None, 'ИТОГО', None, None, len(g), s.План, s.Факт, _pct(s.Факт, s.План),
                s.Витрина, _pct(s.Витрина, s.План), s.Касса, _pct(s.Касса, s.План), None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = Font(bold=True); c.fill = GRN; c.border = BOX
            if j in (9, 11, 13): c.number_format = '0%'
        totals.append(s); r += 1
    if len(days):
        s = days[['План', 'Факт', 'Витрина', 'Касса']].sum(); r += 1
        d0, d1 = days.day.min(), days.day.max()
        vals = [f"01-{d1:%d.%m.%Y}", None, 'ИТОГО', None, None, len(days), s.План, s.Факт, _pct(s.Факт, s.План),
                s.Витрина, _pct(s.Витрина, s.План), s.Касса, _pct(s.Касса, s.План), None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = Font(bold=True); c.fill = GRN; c.border = BOX
            if j in (9, 11, 13): c.number_format = '0%'
    for j, w in enumerate([8, 5, 34, 10, 18, 9, 9, 9, 10, 10, 10, 10, 10, 70], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'

    # ---------- Ежемесячный ----------
    ws = wb.create_sheet('Ежемесячный')
    cols = ['Месяц', '№', 'Менеджер', 'Должность', 'Регион', 'Кол-во дней', 'План', 'Визит факт', '%Выполнения',
            'Витрина Мерч', '% Витрина Мерч', 'Касса Мерч', '% Касса Мерч', 'Шубҳали визит', 'Изоҳ']
    _hdr(ws, 1, cols)
    per = days.groupby('ТП').agg(Регион=('Регион', 'first'), d=('day', 'nunique'), План=('План', 'sum'),
                                 Факт=('Факт', 'sum'), Витрина=('Витрина', 'sum'), Касса=('Касса', 'sum')).reset_index()
    sus = tpx.groupby('ТП')['Шубҳали'].sum()
    flagcnt = {k: tpx.groupby('ТП')['f_' + k].sum() for k in SHORT if 'f_' + k in tpx}
    per = per.sort_values(['Регион', 'ТП'])
    for n, x in enumerate(per.itertuples(index=False), 1):
        note = "; ".join(f"{int(flagcnt[k].get(x.ТП, 0))} та {v}" for k, v in SHORT.items()
                         if k in flagcnt and flagcnt[k].get(x.ТП, 0))
        if tpx.loc[tpx['ТП'] == x.ТП, 'Янги'].any():
            note = "ЯНГИ ТП; " + note
        vals = [MON[days.day.max().month], n, x.ТП, 'ТП', x.Регион, x.d, x.План, x.Факт, _pct(x.Факт, x.План),
                x.Витрина, _pct(x.Витрина, x.План), x.Касса, _pct(x.Касса, x.План), int(sus.get(x.ТП, 0)), note]
        for j, v in enumerate(vals, 1):
            c = ws.cell(n + 1, j, v); c.border = BOX
            if j in (9, 11, 13): c.number_format = '0%'
        if x.Факт and sus.get(x.ТП, 0) / x.Факт >= 0.3:
            ws.cell(n + 1, 14).fill = RED
    rr = len(per) + 2
    s = per[['d', 'План', 'Факт', 'Витрина', 'Касса']].sum()
    vals = ['ИТОГО', None, None, None, None, s.d, s.План, s.Факт, _pct(s.Факт, s.План), s.Витрина, _pct(s.Витрина, s.План),
            s.Касса, _pct(s.Касса, s.План), int(sus.sum()), None]
    for j, v in enumerate(vals, 1):
        c = ws.cell(rr, j, v); c.font = Font(bold=True); c.fill = GRN; c.border = BOX
        if j in (9, 11, 13): c.number_format = '0%'
    # region table on the right
    c0 = 17
    for j, h in enumerate(['Регион', 'План', 'Факт визит', '% визит', 'Витрина Мерч', '%', 'Касса Мерч', '%'], c0):
        x = ws.cell(1, j, h); x.font = HF; x.fill = HDR; x.border = BOX; x.alignment = Alignment(wrap_text=True, horizontal='center')
    reg = per.groupby('Регион')[['План', 'Факт', 'Витрина', 'Касса']].sum().reset_index()
    for i, x in enumerate(reg.itertuples(index=False), 2):
        for j, v in enumerate([x.Регион, x.План, x.Факт, _pct(x.Факт, x.План), x.Витрина, _pct(x.Витрина, x.План),
                               x.Касса, _pct(x.Касса, x.План)], c0):
            c = ws.cell(i, j, v); c.border = BOX
            if j in (c0 + 3, c0 + 5, c0 + 7): c.number_format = '0%'
    i = len(reg) + 2; s = reg[['План', 'Факт', 'Витрина', 'Касса']].sum()
    for j, v in enumerate(['Итого', s.План, s.Факт, _pct(s.Факт, s.План), s.Витрина, _pct(s.Витрина, s.План), s.Касса,
                           _pct(s.Касса, s.План)], c0):
        c = ws.cell(i, j, v); c.font = Font(bold=True); c.fill = GRN; c.border = BOX
        if j in (c0 + 3, c0 + 5, c0 + 7): c.number_format = '0%'
    for j, w in enumerate([10, 5, 34, 10, 18, 9, 9, 9, 10, 10, 10, 10, 10, 10, 60, 2, 18, 9, 9, 9, 10, 9, 10, 9], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'

    # ---------- Визитлар тўлиқ ----------
    ws = wb.create_sheet('Визитлар тўлиқ')
    cols = ['ID', 'Дата', 'Ёпилган вақт', 'ТП', 'Иш ҳудуди', 'Аптека', 'Мерч хулоса', 'Витрина', 'Касса',
            'ИИ изоҳи (ҳар бир фото)', 'Шубҳа / изоҳ', 'Фото 1', 'Фото 2']
    _hdr(ws, 1, cols)
    mc = merch.set_index('id')['comment'] if len(merch) and 'comment' in merch else pd.Series(dtype=str)
    v = tpx.sort_values(['day', 'ТП', 'dt'])
    for i, x in enumerate(v.to_dict('records'), 2):
        cat = x['cat'] if isinstance(x['cat'], str) else None
        if cat:
            mtxt = CAT.get(cat, 'Текширилмади')
        else:
            mtxt = 'қўлда (ИИ дан олдин)' if x['day'] < pd.Timestamp('2026-10-26') else 'Текширилмади'
        vals = [x['id'], x['Дата'], x['Время завершения визита'], x['ТП'], x['Иш ҳудуди'], x['Контр-агент/Врач'], mtxt,
                ('Да' if x['v'] else 'Нет') if cat else '', ('Да' if x['k'] else 'Нет') if cat else '',
                mc.get(x['id'], ''), x['Изоҳ'],
                x['Фото 1'], x['Фото 2']]
        for j, val in enumerate(vals, 1):
            c = ws.cell(i, j, val)
            if j in (12, 13) and isinstance(val, str) and val.startswith('http'):
                c.hyperlink = val; c.value = 'Очиш'; c.font = Font(color='0563C1', underline='single')
        if x['Шубҳали']:
            ws.cell(i, 11).fill = YEL
        if cat == 'merch':
            ws.cell(i, 7).fill = PatternFill('solid', fgColor='C6EFCE')
        elif cat in ('none', 'suspicious'):
            ws.cell(i, 7).fill = RED if cat == 'none' else YEL
    for j, w in enumerate([9, 11, 16, 32, 16, 40, 18, 8, 8, 70, 80, 8, 8], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

    # ---------- Аптека муаммолари ----------
    ws = wb.create_sheet('Аптека муаммолари')
    cols = ['Аптека', 'Муаммо', 'ТП (лар)', 'Визитлар сони', 'Изоҳ']
    _hdr(ws, 1, cols)
    rows = []
    for k, title in [('region_out', 'Бошқа ҳудуддаги аптека'), ('cluster', 'Бир нуқтада 5+ аптека'),
                     ('multi', '3+ ТП борган'), ('over7', '7+ визит'), ('loc', 'Локацияси ўзгарган')]:
        if 'f_' + k not in tpx:
            continue
        for obj, g in tpx[tpx['f_' + k]].groupby('Контр-агент/Врач'):
            note = g['Изоҳ'].str.split('; ').explode()
            note = note[note.str.contains({'region_out': 'координатаси', 'cluster': 'бир нуқтада', 'multi': 'та ТП борган',
                                           'over7': 'марта визит', 'loc': 'локацияси ўзгартирилган'}[k])]
            rows.append([obj, title, ", ".join(sorted(g['ТП'].unique())), len(g), note.iloc[0] if len(note) else ''])
    for i, row in enumerate(sorted(rows, key=lambda r: (r[1], -r[3])), 2):
        for j, val in enumerate(row, 1):
            ws.cell(i, j, val)
    for j, w in enumerate([45, 24, 50, 12, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

    # ---------- Локация ўзгаришлари ----------
    if loc_changes is not None and len(loc_changes):
        ws = wb.create_sheet('Локация ўзгаришлари')
        _hdr(ws, 1, list(loc_changes.columns))
        for i, row in enumerate(loc_changes.itertuples(index=False), 2):
            for j, val in enumerate(row, 1):
                ws.cell(i, j, val)
        for j in range(1, len(loc_changes.columns) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 20
        ws.column_dimensions['D'].width = 40; ws.column_dimensions['G'].width = 45

    # ---------- Изоҳ ----------
    ws = wb.create_sheet('Изоҳ')
    lines = [month_title, '',
             'План: ТП шу куни камида 1 та визит қилган бўлса 1 кунлик план қўйилади (10 ёки 15, рўйхат бўйича; янги ТП — 15).',
             'Визит факт: барча визитлар саналади, шубҳалилари ҳам — улар «Комментария» ва «Визитлар тўлиқ» да кўрсатилади.',
             'Витрина / Касса Мерч: ҳар бир визит фотоси ИИ (бепул Google Gemini) да текширилади. Қоида: Риномакс ХОТ ва Форсил',
             '  иккаласи БИТТА расмда; Витрина — витринада/полкада, Касса — касса олдида. Яқин план (аптека кўринмайди) — мерчга қўшилмайди.',
             '  26.10.2026 гача бўлган кунлар учун мерч сизнинг файлингиздан (қўлда киритилган) олинган.',
             'Шубҳа белгилари: тез ёпилган (≤1 дақ), фото йўқ/бир хил/қайта ишлатилган, такрорий визит, иш вақтидан ташқари,',
             '  аптека локацияси ўзгартирилган (ҳар соатлик кузатув), аптека координатаси ТП ишлайдиган ҳудуддан ташқарида,',
             '  бир нуқтада 5+ аптека, бир ойда аптекага 3+ ТП, бир ТП бир аптекага ойда 7 тадан кўп.',
             'Ҳар бир ҳолат хотирада сақланади; такрорланса изоҳда «такрорий» деб олдинги саналари билан ёзилади.']
    for t in lines:
        ws.append([t])
    ws.column_dimensions['A'].width = 150
    wb.move_sheet('Изоҳ', offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    wb.save(path)
    return days, tpx
